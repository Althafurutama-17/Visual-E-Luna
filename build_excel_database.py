import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

def parse_dmy(dmy_str):
    """Konversi string DD/MM/YYYY ke objek datetime.date untuk Excel native date"""
    if not dmy_str:
        return None
    p = dmy_str.split('/')
    if len(p) == 3:
        return datetime.date(int(p[2]), int(p[1]), int(p[0]))
    return None

def build_bpbd_excel_database(output_path):
    wb = Workbook()
    
    # -------------------------------------------------------------
    # PALET WARNA RESMI & STYLING (E-LUNA & BPBD JAWA BARAT)
    # -------------------------------------------------------------
    FONT_FAMILY = "Segoe UI"
    
    font_header = Font(name=FONT_FAMILY, size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="0E63AA", end_color="0E63AA", fill_type="solid") # Biru E-Luna
    
    font_sub_header = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
    fill_sub_header = PatternFill(start_color="11538C", end_color="11538C", fill_type="solid")
    
    font_data = Font(name=FONT_FAMILY, size=10, color="1E293B")
    font_bold = Font(name=FONT_FAMILY, size=10, bold=True, color="1E293B")
    font_code = Font(name="Consolas", size=9, bold=True, color="0E63AA")
    
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_total = PatternFill(start_color="EEF2F6", end_color="EEF2F6", fill_type="solid")
    
    # Border
    thin_border_side = Side(style='thin', color="CBD5E1")
    double_bottom_side = Side(style='double', color="0E63AA")
    thick_top_side = Side(style='thin', color="0E63AA")
    
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_total = Border(top=thick_top_side, bottom=double_bottom_side, left=thin_border_side, right=thin_border_side)
    
    align_left = Alignment(horizontal='left', vertical='center')
    align_center = Alignment(horizontal='center', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')
    align_header = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # -------------------------------------------------------------
    # 1. SHEET _RefData (Master Dropdown References)
    # -------------------------------------------------------------
    ws_ref = wb.active
    ws_ref.title = "_RefData"
    ws_ref.views.sheetView[0].showGridLines = True
    
    # Kolom A: Sumber Pendanaan (Tepat 3 Sesuai Instruksi User)
    sumber_dana = ["APBN", "APBD", "CSR"]
    ws_ref['A1'] = "Sumber Pendanaan"
    for r_idx, val in enumerate(sumber_dana, start=2):
        ws_ref[f'A{r_idx}'] = val
        
    # Kolom B: Kategori Barang (4 Kategori Resmi)
    kategori_barang = ["Sandang", "Pangan", "Papan", "Logistik Lainnya"]
    ws_ref['B1'] = "Kategori Barang"
    for r_idx, val in enumerate(kategori_barang, start=2):
        ws_ref[f'B{r_idx}'] = val
        
    # Kolom C: 27 Kab/Kota Jawa Barat
    kab_kota_jabar = [
        "-",
        "Kabupaten Bandung",
        "Kabupaten Bandung Barat",
        "Kabupaten Bekasi",
        "Kabupaten Bogor",
        "Kabupaten Ciamis",
        "Kabupaten Cianjur",
        "Kabupaten Cirebon",
        "Kabupaten Garut",
        "Kabupaten Indramayu",
        "Kabupaten Karawang",
        "Kabupaten Kuningan",
        "Kabupaten Majalengka",
        "Kabupaten Pangandaran",
        "Kabupaten Purwakarta",
        "Kabupaten Subang",
        "Kabupaten Sukabumi",
        "Kabupaten Sumedang",
        "Kabupaten Tasikmalaya",
        "Kota Bandung",
        "Kota Banjar",
        "Kota Bekasi",
        "Kota Bogor",
        "Kota Cimahi",
        "Kota Cirebon",
        "Kota Depok",
        "Kota Sukabumi",
        "Kota Tasikmalaya"
    ]
    ws_ref['C1'] = "Kab/Kota Tujuan"
    for r_idx, val in enumerate(kab_kota_jabar, start=2):
        ws_ref[f'C{r_idx}'] = val

    for col in ['A', 'B', 'C']:
        cell = ws_ref[f'{col}1']
        cell.font = font_sub_header
        cell.fill = fill_sub_header
        cell.alignment = align_center

    # -------------------------------------------------------------
    # 2. MASTER DATASET MUTASI LOGISTIK (REALISTIS BPBD JAWA BARAT)
    # -------------------------------------------------------------
    raw_mutasi = [
        # Transaksi Penerimaan (BAST Masuk)
        {"dana": "APBN", "barang": "Beras Premium 5kg", "kat": "Pangan", "in": 12000, "out": 0, "tujuan": "-", "exp": "15/04/2027", "tgl_in": "01/09/2026", "tgl_out": "", "doc": "BAST-2026-IX-0012"},
        {"dana": "APBD", "barang": "Makanan Siap Saji (Paket 4 Lauk)", "kat": "Pangan", "in": 3500, "out": 0, "tujuan": "-", "exp": "18/10/2026", "tgl_in": "02/09/2026", "tgl_out": "", "doc": "BAST-2026-IX-0015"},
        {"dana": "CSR", "barang": "Biskuit Balita Tambahan Gizi", "kat": "Pangan", "in": 1500, "out": 0, "tujuan": "-", "exp": "20/09/2026", "tgl_in": "05/09/2026", "tgl_out": "", "doc": "BAST-2026-IX-0018"},
        {"dana": "APBN", "barang": "Paket Sandang Dewasa", "kat": "Sandang", "in": 2500, "out": 0, "tujuan": "-", "exp": "", "tgl_in": "06/09/2026", "tgl_out": "", "doc": "BAST-2026-IX-0021"},
        {"dana": "APBD", "barang": "Matras & Selimut Wol", "kat": "Sandang", "in": 1800, "out": 0, "tujuan": "-", "exp": "", "tgl_in": "08/09/2026", "tgl_out": "", "doc": "BAST-2026-IX-0024"},
        {"dana": "APBN", "barang": "Tenda Pengungsi 4x4m", "kat": "Papan", "in": 120, "out": 0, "tujuan": "-", "exp": "", "tgl_in": "10/09/2026", "tgl_out": "", "doc": "BAST-2026-IX-0027"},
        {"dana": "APBD", "barang": "Terpal Gulung Hunian Darurat", "kat": "Papan", "in": 800, "out": 0, "tujuan": "-", "exp": "", "tgl_in": "12/09/2026", "tgl_out": "", "doc": "BAST-2026-IX-0030"},
        {"dana": "APBN", "barang": "Perahu Karet Evakuasi + Mopel", "kat": "Logistik Lainnya", "in": 15, "out": 0, "tujuan": "-", "exp": "", "tgl_in": "15/09/2026", "tgl_out": "", "doc": "BAST-2026-IX-0033"},
        {"dana": "APBD", "barang": "Cairan Antiseptik Sanitasi 1L", "kat": "Logistik Lainnya", "in": 600, "out": 0, "tujuan": "-", "exp": "30/08/2026", "tgl_in": "16/09/2026", "tgl_out": "", "doc": "BAST-2026-IX-0036"},
        {"dana": "CSR", "barang": "Susu Formula Balita Tanggap Darurat", "kat": "Pangan", "in": 800, "out": 0, "tujuan": "-", "exp": "25/10/2026", "tgl_in": "18/09/2026", "tgl_out": "", "doc": "BAST-2026-IX-0039"},
        {"dana": "APBN", "barang": "Lauk Pauk Kaleng Saus Tomat", "kat": "Pangan", "in": 4000, "out": 0, "tujuan": "-", "exp": "02/11/2026", "tgl_in": "20/09/2026", "tgl_out": "", "doc": "BAST-2026-IX-0042"},
        {"dana": "APBD", "barang": "Family Kit Sanitasi Bencana", "kat": "Logistik Lainnya", "in": 950, "out": 0, "tujuan": "-", "exp": "", "tgl_in": "22/09/2026", "tgl_out": "", "doc": "BAST-2026-IX-0045"},

        # Transaksi Pengeluaran / Distribusi (Surat Jalan Keluar)
        {"dana": "APBN", "barang": "Beras Premium 5kg", "kat": "Pangan", "in": 0, "out": 2500, "tujuan": "Kabupaten Bogor", "exp": "", "tgl_in": "", "tgl_out": "15/09/2026", "doc": "SJ-2026-IX-0081"},
        {"dana": "APBD", "barang": "Makanan Siap Saji (Paket 4 Lauk)", "kat": "Pangan", "in": 0, "out": 1200, "tujuan": "Kabupaten Cianjur", "exp": "", "tgl_in": "", "tgl_out": "16/09/2026", "doc": "SJ-2026-IX-0083"},
        {"dana": "APBN", "barang": "Paket Sandang Dewasa", "kat": "Sandang", "in": 0, "out": 600, "tujuan": "Kabupaten Garut", "exp": "", "tgl_in": "", "tgl_out": "18/09/2026", "doc": "SJ-2026-IX-0085"},
        {"dana": "APBD", "barang": "Matras & Selimut Wol", "kat": "Sandang", "in": 0, "out": 450, "tujuan": "Kabupaten Bandung", "exp": "", "tgl_in": "", "tgl_out": "20/09/2026", "doc": "SJ-2026-IX-0088"},
        {"dana": "APBN", "barang": "Tenda Pengungsi 4x4m", "kat": "Papan", "in": 0, "out": 25, "tujuan": "Kabupaten Sukabumi", "exp": "", "tgl_in": "", "tgl_out": "22/09/2026", "doc": "SJ-2026-IX-0090"},
        {"dana": "APBD", "barang": "Terpal Gulung Hunian Darurat", "kat": "Papan", "in": 0, "out": 200, "tujuan": "Kabupaten Ciamis", "exp": "", "tgl_in": "", "tgl_out": "23/09/2026", "doc": "SJ-2026-IX-0092"},
        {"dana": "APBN", "barang": "Perahu Karet Evakuasi + Mopel", "kat": "Logistik Lainnya", "in": 0, "out": 3, "tujuan": "Kabupaten Bekasi", "exp": "", "tgl_in": "", "tgl_out": "25/09/2026", "doc": "SJ-2026-IX-0095"},
        {"dana": "CSR", "barang": "Biskuit Balita Tambahan Gizi", "kat": "Pangan", "in": 0, "out": 300, "tujuan": "Kota Bandung", "exp": "", "tgl_in": "", "tgl_out": "26/09/2026", "doc": "SJ-2026-IX-0098"},
        {"dana": "APBD", "barang": "Family Kit Sanitasi Bencana", "kat": "Logistik Lainnya", "in": 0, "out": 180, "tujuan": "Kabupaten Subang", "exp": "", "tgl_in": "", "tgl_out": "27/09/2026", "doc": "SJ-2026-IX-0101"},
        {"dana": "APBN", "barang": "Lauk Pauk Kaleng Saus Tomat", "kat": "Pangan", "in": 0, "out": 850, "tujuan": "Kota Bekasi", "exp": "", "tgl_in": "", "tgl_out": "28/09/2026", "doc": "SJ-2026-IX-0104"},
        {"dana": "CSR", "barang": "Susu Formula Balita Tanggap Darurat", "kat": "Pangan", "in": 0, "out": 150, "tujuan": "Kabupaten Purwakarta", "exp": "", "tgl_in": "", "tgl_out": "29/09/2026", "doc": "SJ-2026-IX-0107"},
        {"dana": "APBD", "barang": "Beras Premium 5kg", "kat": "Pangan", "in": 0, "out": 1500, "tujuan": "Kota Bogor", "exp": "", "tgl_in": "", "tgl_out": "01/10/2026", "doc": "SJ-2026-X-0110"}
    ]

    # Data Validasi Dropdown
    dv_dana = DataValidation(type="list", formula1="'_RefData'!$A$2:$A$4", allow_blank=True)
    dv_kat = DataValidation(type="list", formula1="'_RefData'!$B$2:$B$5", allow_blank=True)
    dv_tujuan = DataValidation(type="list", formula1="'_RefData'!$C$2:$C$29", allow_blank=True)

    # -------------------------------------------------------------
    # 3. SHEET 1: List Data (Tabel: tblListData)
    # -------------------------------------------------------------
    ws_list = wb.create_sheet(title="List Data")
    ws_list.views.sheetView[0].showGridLines = True
    ws_list.freeze_panes = "A2"
    ws_list.add_data_validation(dv_dana)
    ws_list.add_data_validation(dv_kat)
    ws_list.add_data_validation(dv_tujuan)

    headers_list = [
        "Sumber Pendanaan",
        "Barang",
        "Kategori Barang",
        "Jumlah Masuk",
        "Jumlah Keluar",
        "Kab/Kota Tujuan",
        "Tanggal Kadaluwarsa",
        "Tanggal barang diterima",
        "Tanggal barang dikeluarkan",
        "Nomor Dokumen"
    ]
    ws_list.append(headers_list)
    ws_list.row_dimensions[1].height = 26

    for r_idx, item in enumerate(raw_mutasi, start=2):
        row_vals = [
            item["dana"],
            item["barang"],
            item["kat"],
            item["in"],
            item["out"],
            item["tujuan"],
            parse_dmy(item["exp"]),
            parse_dmy(item["tgl_in"]),
            parse_dmy(item["tgl_out"]),
            item["doc"]
        ]
        ws_list.append(row_vals)
        ws_list.row_dimensions[r_idx].height = 20

    max_row_list = len(raw_mutasi) + 1
    for col_idx in range(1, 11):
        cell_header = ws_list.cell(row=1, column=col_idx)
        cell_header.font = font_header
        cell_header.fill = fill_header
        cell_header.alignment = align_header
        cell_header.border = border_cell

    for row_idx in range(2, max_row_list + 1):
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, 11):
            c = ws_list.cell(row=row_idx, column=col_idx)
            c.font = font_data
            c.border = border_cell
            if not is_even:
                c.fill = fill_zebra
            
            if col_idx in [1, 3]:
                c.alignment = align_center
            elif col_idx in [4, 5]:
                c.number_format = '#,##0'
                c.alignment = align_right
                if c.value and c.value > 0:
                    c.font = font_bold
            elif col_idx == 6:
                c.alignment = align_center
            elif col_idx in [7, 8, 9]:
                c.alignment = align_center
                if c.value is not None:
                    c.number_format = 'DD/MM/YYYY'
            elif col_idx == 10:
                c.font = font_code
                c.alignment = align_center

    dv_dana.add(f"A2:A{max_row_list}")
    dv_kat.add(f"C2:C{max_row_list}")
    dv_tujuan.add(f"F2:F{max_row_list}")

    tab_list = Table(displayName="tblListData", ref=f"A1:J{max_row_list}")
    tab_list.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws_list.add_table(tab_list)

    # -------------------------------------------------------------
    # 4. SHEET 2: Penerimaan (Tabel: tblPenerimaan)
    # -------------------------------------------------------------
    ws_penerimaan = wb.create_sheet(title="Penerimaan")
    ws_penerimaan.views.sheetView[0].showGridLines = True
    ws_penerimaan.freeze_panes = "A2"
    ws_penerimaan.add_data_validation(dv_dana)
    ws_penerimaan.add_data_validation(dv_kat)

    headers_penerimaan = [
        "Sumber Pendanaan",
        "Barang",
        "Kategori Barang",
        "Jumlah Masuk",
        "Tanggal Kadaluwarsa",
        "Tanggal barang diterima",
        "Nomor Dokumen"
    ]
    ws_penerimaan.append(headers_penerimaan)
    ws_penerimaan.row_dimensions[1].height = 26

    penerimaan_data = [item for item in raw_mutasi if item["in"] > 0]
    for r_idx, item in enumerate(penerimaan_data, start=2):
        row_vals = [
            item["dana"],
            item["barang"],
            item["kat"],
            item["in"],
            parse_dmy(item["exp"]),
            parse_dmy(item["tgl_in"]),
            item["doc"]
        ]
        ws_penerimaan.append(row_vals)
        ws_penerimaan.row_dimensions[r_idx].height = 20

    max_row_penerimaan = len(penerimaan_data) + 1
    for col_idx in range(1, 8):
        cell_header = ws_penerimaan.cell(row=1, column=col_idx)
        cell_header.font = font_header
        cell_header.fill = fill_header
        cell_header.alignment = align_header
        cell_header.border = border_cell

    for row_idx in range(2, max_row_penerimaan + 1):
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, 8):
            c = ws_penerimaan.cell(row=row_idx, column=col_idx)
            c.font = font_data
            c.border = border_cell
            if not is_even:
                c.fill = fill_zebra
            if col_idx in [1, 3]:
                c.alignment = align_center
            elif col_idx == 4:
                c.number_format = '#,##0'
                c.alignment = align_right
                c.font = font_bold
            elif col_idx in [5, 6]:
                c.alignment = align_center
                if c.value is not None:
                    c.number_format = 'DD/MM/YYYY'
            elif col_idx == 7:
                c.font = font_code
                c.alignment = align_center

    dv_dana.add(f"A2:A{max_row_penerimaan}")
    dv_kat.add(f"C2:C{max_row_penerimaan}")

    tab_penerimaan = Table(displayName="tblPenerimaan", ref=f"A1:G{max_row_penerimaan}")
    tab_penerimaan.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws_penerimaan.add_table(tab_penerimaan)

    # Baris Total di Penerimaan
    tot_row_p = max_row_penerimaan + 1
    ws_penerimaan.cell(row=tot_row_p, column=1, value="TOTAL PENERIMAAN").font = font_bold
    ws_penerimaan.cell(row=tot_row_p, column=1).alignment = align_center
    ws_penerimaan.cell(row=tot_row_p, column=4, value=f"=SUM(D2:D{max_row_penerimaan})").number_format = '#,##0'
    ws_penerimaan.cell(row=tot_row_p, column=4).font = font_bold
    ws_penerimaan.cell(row=tot_row_p, column=4).alignment = align_right
    for c_idx in range(1, 8):
        cell_t = ws_penerimaan.cell(row=tot_row_p, column=c_idx)
        cell_t.fill = fill_total
        cell_t.border = border_total

    # -------------------------------------------------------------
    # 5. SHEET 3: Distribusi (Tabel: tblDistribusi)
    # -------------------------------------------------------------
    ws_distribusi = wb.create_sheet(title="Distribusi")
    ws_distribusi.views.sheetView[0].showGridLines = True
    ws_distribusi.freeze_panes = "A2"
    ws_distribusi.add_data_validation(dv_dana)
    ws_distribusi.add_data_validation(dv_kat)
    ws_distribusi.add_data_validation(dv_tujuan)

    headers_distribusi = [
        "Sumber Pendanaan",
        "Barang",
        "Kategori Barang",
        "Jumlah Keluar",
        "Kab/Kota Tujuan",
        "Tanggal barang dikeluarkan",
        "Nomor Dokumen"
    ]
    ws_distribusi.append(headers_distribusi)
    ws_distribusi.row_dimensions[1].height = 26

    distribusi_data = [item for item in raw_mutasi if item["out"] > 0]
    for r_idx, item in enumerate(distribusi_data, start=2):
        row_vals = [
            item["dana"],
            item["barang"],
            item["kat"],
            item["out"],
            item["tujuan"],
            parse_dmy(item["tgl_out"]),
            item["doc"]
        ]
        ws_distribusi.append(row_vals)
        ws_distribusi.row_dimensions[r_idx].height = 20

    max_row_distribusi = len(distribusi_data) + 1
    for col_idx in range(1, 8):
        cell_header = ws_distribusi.cell(row=1, column=col_idx)
        cell_header.font = font_header
        cell_header.fill = fill_header
        cell_header.alignment = align_header
        cell_header.border = border_cell

    for row_idx in range(2, max_row_distribusi + 1):
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, 8):
            c = ws_distribusi.cell(row=row_idx, column=col_idx)
            c.font = font_data
            c.border = border_cell
            if not is_even:
                c.fill = fill_zebra
            if col_idx in [1, 3, 5]:
                c.alignment = align_center
            elif col_idx == 4:
                c.number_format = '#,##0'
                c.alignment = align_right
                c.font = font_bold
            elif col_idx == 6:
                c.alignment = align_center
                if c.value is not None:
                    c.number_format = 'DD/MM/YYYY'
            elif col_idx == 7:
                c.font = font_code
                c.alignment = align_center

    dv_dana.add(f"A2:A{max_row_distribusi}")
    dv_kat.add(f"C2:C{max_row_distribusi}")
    dv_tujuan.add(f"E2:E{max_row_distribusi}")

    tab_distribusi = Table(displayName="tblDistribusi", ref=f"A1:G{max_row_distribusi}")
    tab_distribusi.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws_distribusi.add_table(tab_distribusi)

    tot_row_d = max_row_distribusi + 1
    ws_distribusi.cell(row=tot_row_d, column=1, value="TOTAL DISTRIBUSI").font = font_bold
    ws_distribusi.cell(row=tot_row_d, column=1).alignment = align_center
    ws_distribusi.cell(row=tot_row_d, column=4, value=f"=SUM(D2:D{max_row_distribusi})").number_format = '#,##0'
    ws_distribusi.cell(row=tot_row_d, column=4).font = font_bold
    ws_distribusi.cell(row=tot_row_d, column=4).alignment = align_right
    for c_idx in range(1, 8):
        cell_t = ws_distribusi.cell(row=tot_row_d, column=c_idx)
        cell_t.fill = fill_total
        cell_t.border = border_total

    # -------------------------------------------------------------
    # 6. SHEET 4: Stok Awal (Sheet Bantu / Master Saldo Lalu)
    # -------------------------------------------------------------
    ws_stok_awal = wb.create_sheet(title="Stok Awal")
    ws_stok_awal.views.sheetView[0].showGridLines = True
    ws_stok_awal.freeze_panes = "A2"
    ws_stok_awal.add_data_validation(dv_dana)
    ws_stok_awal.add_data_validation(dv_kat)

    headers_stok_awal = [
        "Kunci",
        "Sumber Pendanaan",
        "Barang",
        "Kategori Barang",
        "Stok Akhir Periode Sebelumnya"
    ]
    ws_stok_awal.append(headers_stok_awal)
    ws_stok_awal.row_dimensions[1].height = 26

    master_komoditas = [
        {"dana": "APBN", "barang": "Beras Premium 5kg", "kat": "Pangan", "awal": 40000},
        {"dana": "APBD", "barang": "Beras Premium 5kg", "kat": "Pangan", "awal": 18000},
        {"dana": "APBD", "barang": "Makanan Siap Saji (Paket 4 Lauk)", "kat": "Pangan", "awal": 4500},
        {"dana": "CSR", "barang": "Biskuit Balita Tambahan Gizi", "kat": "Pangan", "awal": 850},
        {"dana": "CSR", "barang": "Susu Formula Balita Tanggap Darurat", "kat": "Pangan", "awal": 500},
        {"dana": "APBN", "barang": "Lauk Pauk Kaleng Saus Tomat", "kat": "Pangan", "awal": 3200},
        {"dana": "APBN", "barang": "Paket Sandang Dewasa", "kat": "Sandang", "awal": 8500},
        {"dana": "APBD", "barang": "Matras & Selimut Wol", "kat": "Sandang", "awal": 6000},
        {"dana": "APBN", "barang": "Tenda Pengungsi 4x4m", "kat": "Papan", "awal": 450},
        {"dana": "APBD", "barang": "Terpal Gulung Hunian Darurat", "kat": "Papan", "awal": 8750},
        {"dana": "APBN", "barang": "Perahu Karet Evakuasi + Mopel", "kat": "Logistik Lainnya", "awal": 45},
        {"dana": "APBD", "barang": "Cairan Antiseptik Sanitasi 1L", "kat": "Logistik Lainnya", "awal": 1200},
        {"dana": "APBD", "barang": "Family Kit Sanitasi Bencana", "kat": "Logistik Lainnya", "awal": 15555}
    ]

    for r_idx, km in enumerate(master_komoditas, start=2):
        formula_kunci = f'=B{r_idx}&"|"&C{r_idx}&"|"&D{r_idx}'
        row_vals = [
            formula_kunci,
            km["dana"],
            km["barang"],
            km["kat"],
            km["awal"]
        ]
        ws_stok_awal.append(row_vals)
        ws_stok_awal.row_dimensions[r_idx].height = 20

    max_row_awal = len(master_komoditas) + 1
    for col_idx in range(1, 6):
        cell_header = ws_stok_awal.cell(row=1, column=col_idx)
        cell_header.font = font_header
        cell_header.fill = fill_header
        cell_header.alignment = align_header
        cell_header.border = border_cell

    for row_idx in range(2, max_row_awal + 1):
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, 6):
            c = ws_stok_awal.cell(row=row_idx, column=col_idx)
            c.font = font_data
            c.border = border_cell
            if not is_even:
                c.fill = fill_zebra
            if col_idx == 1:
                c.font = font_code
                c.alignment = align_center
            elif col_idx in [2, 4]:
                c.alignment = align_center
            elif col_idx == 5:
                c.number_format = '#,##0'
                c.alignment = align_right
                c.font = font_bold

    dv_dana.add(f"B2:B{max_row_awal}")
    dv_kat.add(f"D2:D{max_row_awal}")

    tab_awal = Table(displayName="tblStokAwal", ref=f"A1:E{max_row_awal}")
    tab_awal.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws_stok_awal.add_table(tab_awal)

    # -------------------------------------------------------------
    # 7. SHEET 5: Stok (Tabel: tblStok)
    # Perhitungan Formula Dinamis: Stok Awal + Masuk - Keluar
    # -------------------------------------------------------------
    ws_stok = wb.create_sheet(title="Stok")
    ws_stok.views.sheetView[0].showGridLines = True
    ws_stok.freeze_panes = "A2"
    ws_stok.add_data_validation(dv_dana)
    ws_stok.add_data_validation(dv_kat)

    headers_stok = [
        "Sumber Pendanaan",
        "Barang",
        "Kategori Barang",
        "Stok"
    ]
    ws_stok.append(headers_stok)
    ws_stok.row_dimensions[1].height = 26

    for r_idx, km in enumerate(master_komoditas, start=2):
        # Formula:
        # IFERROR(VLOOKUP(kunci, 'Stok Awal'!$A:$E, 5, FALSE), 0) + SUMIFS(Penerimaan) - SUMIFS(Distribusi)
        formula_stok = (
            f'=IFERROR(VLOOKUP(A{r_idx}&"|"&B{r_idx}&"|"&C{r_idx}, \'Stok Awal\'!$A:$E, 5, FALSE), 0)'
            f'+SUMIFS(Penerimaan!$D:$D, Penerimaan!$A:$A, A{r_idx}, Penerimaan!$B:$B, B{r_idx}, Penerimaan!$C:$C, C{r_idx})'
            f'-SUMIFS(Distribusi!$D:$D, Distribusi!$A:$A, A{r_idx}, Distribusi!$B:$B, B{r_idx}, Distribusi!$C:$C, C{r_idx})'
        )
        row_vals = [
            km["dana"],
            km["barang"],
            km["kat"],
            formula_stok
        ]
        ws_stok.append(row_vals)
        ws_stok.row_dimensions[r_idx].height = 20

    max_row_stok = len(master_komoditas) + 1
    for col_idx in range(1, 5):
        cell_header = ws_stok.cell(row=1, column=col_idx)
        cell_header.font = font_header
        cell_header.fill = fill_header
        cell_header.alignment = align_header
        cell_header.border = border_cell

    for row_idx in range(2, max_row_stok + 1):
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, 5):
            c = ws_stok.cell(row=row_idx, column=col_idx)
            c.font = font_data
            c.border = border_cell
            if not is_even:
                c.fill = fill_zebra
            if col_idx in [1, 3]:
                c.alignment = align_center
            elif col_idx == 4:
                c.number_format = '#,##0'
                c.alignment = align_right
                c.font = font_bold

    dv_dana.add(f"A2:A{max_row_stok}")
    dv_kat.add(f"C2:C{max_row_stok}")

    tab_stok = Table(displayName="tblStok", ref=f"A1:D{max_row_stok}")
    tab_stok.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws_stok.add_table(tab_stok)

    # Total Row Stok
    tot_row_s = max_row_stok + 1
    ws_stok.cell(row=tot_row_s, column=1, value="TOTAL KESELURUHAN STOK").font = font_bold
    ws_stok.cell(row=tot_row_s, column=1).alignment = align_center
    ws_stok.cell(row=tot_row_s, column=4, value=f"=SUM(D2:D{max_row_stok})").number_format = '#,##0'
    ws_stok.cell(row=tot_row_s, column=4).font = font_bold
    ws_stok.cell(row=tot_row_s, column=4).alignment = align_right
    for c_idx in range(1, 5):
        cell_t = ws_stok.cell(row=tot_row_s, column=c_idx)
        cell_t.fill = fill_total
        cell_t.border = border_total

    # -------------------------------------------------------------
    # 8. SHEET 6: Kadaluwarsa (Tabel: tblKadaluwarsa)
    # Formula: Sisa Hari = Tanggal Kadaluwarsa - TODAY()
    # -------------------------------------------------------------
    ws_kadaluwarsa = wb.create_sheet(title="Kadaluwarsa")
    ws_kadaluwarsa.views.sheetView[0].showGridLines = True
    ws_kadaluwarsa.freeze_panes = "A2"
    ws_kadaluwarsa.add_data_validation(dv_dana)
    ws_kadaluwarsa.add_data_validation(dv_kat)

    headers_kadaluwarsa = [
        "Sumber Pendanaan",
        "Barang",
        "Kategori Barang",
        "Nomor Dokumen",
        "Tanggal Kadaluwarsa",
        "Sisa hari"
    ]
    ws_kadaluwarsa.append(headers_kadaluwarsa)
    ws_kadaluwarsa.row_dimensions[1].height = 26

    # Item yang memiliki tanggal kadaluwarsa dari penerimaan
    exp_items = [item for item in raw_mutasi if item["exp"]]
    for r_idx, item in enumerate(exp_items, start=2):
        # Menggunakan formula langsung =E{r_idx}-TODAY() dengan E adalah date object
        formula_sisa = f'=E{r_idx}-TODAY()'
        row_vals = [
            item["dana"],
            item["barang"],
            item["kat"],
            item["doc"],
            parse_dmy(item["exp"]),
            formula_sisa
        ]
        ws_kadaluwarsa.append(row_vals)
        ws_kadaluwarsa.row_dimensions[r_idx].height = 20

    max_row_exp = len(exp_items) + 1
    for col_idx in range(1, 7):
        cell_header = ws_kadaluwarsa.cell(row=1, column=col_idx)
        cell_header.font = font_header
        cell_header.fill = fill_header
        cell_header.alignment = align_header
        cell_header.border = border_cell

    for row_idx in range(2, max_row_exp + 1):
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, 7):
            c = ws_kadaluwarsa.cell(row=row_idx, column=col_idx)
            c.font = font_data
            c.border = border_cell
            if not is_even:
                c.fill = fill_zebra
            if col_idx in [1, 3]:
                c.alignment = align_center
            elif col_idx == 4:
                c.font = font_code
                c.alignment = align_center
            elif col_idx == 5:
                c.alignment = align_center
                if c.value is not None:
                    c.number_format = 'DD/MM/YYYY'
            elif col_idx == 6:
                c.number_format = '#,##0'
                c.alignment = align_center
                c.font = font_bold

    dv_dana.add(f"A2:A{max_row_exp}")
    dv_kat.add(f"C2:C{max_row_exp}")

    tab_exp = Table(displayName="tblKadaluwarsa", ref=f"A1:F{max_row_exp}")
    tab_exp.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws_kadaluwarsa.add_table(tab_exp)

    # Conditional Formatting Kolom F (Sisa Hari)
    # 1. KADALUWARSA (Sisa Hari <= 0) -> Merah
    fill_red = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    font_red = Font(name=FONT_FAMILY, size=10, bold=True, color="991B1B")
    rule_expired = CellIsRule(operator='lessThanOrEqual', formula=['0'], stopIfTrue=True, fill=fill_red, font=font_red)
    ws_kadaluwarsa.conditional_formatting.add(f"F2:F{max_row_exp}", rule_expired)

    # 2. WARNING (1 <= Sisa Hari <= 30) -> Kuning
    fill_yellow = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    font_yellow = Font(name=FONT_FAMILY, size=10, bold=True, color="92400E")
    rule_warning = CellIsRule(operator='between', formula=['1', '30'], stopIfTrue=True, fill=fill_yellow, font=font_yellow)
    ws_kadaluwarsa.conditional_formatting.add(f"F2:F{max_row_exp}", rule_warning)

    # 3. AMAN (Sisa Hari > 30) -> Hijau
    fill_green = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    font_green = Font(name=FONT_FAMILY, size=10, bold=True, color="166534")
    rule_safe = CellIsRule(operator='greaterThan', formula=['30'], stopIfTrue=True, fill=fill_green, font=font_green)
    ws_kadaluwarsa.conditional_formatting.add(f"F2:F{max_row_exp}", rule_safe)

    # -------------------------------------------------------------
    # 9. URUTAN SHEET RESMI & AKTIVASI
    # -------------------------------------------------------------
    # Urutan: List Data, Penerimaan, Distribusi, Stok Awal, Stok, Kadaluwarsa, _RefData
    wb._sheets = [ws_list, ws_penerimaan, ws_distribusi, ws_stok_awal, ws_stok, ws_kadaluwarsa, ws_ref]
    wb.active = ws_list # Default tampilan membuka List Data

    # -------------------------------------------------------------
    # 10. AUTO-FIT COLUMN WIDTHS DI SELURUH WORKSHEET
    # -------------------------------------------------------------
    for ws in [ws_list, ws_penerimaan, ws_distribusi, ws_stok_awal, ws_stok, ws_kadaluwarsa, ws_ref]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if val.startswith('='):
                    val = "FormulaValue"
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # Optimasi lebar kolom khusus
    ws_list.column_dimensions['A'].width = 20
    ws_list.column_dimensions['B'].width = 36
    ws_list.column_dimensions['C'].width = 18
    ws_list.column_dimensions['D'].width = 16
    ws_list.column_dimensions['E'].width = 16
    ws_list.column_dimensions['F'].width = 26
    ws_list.column_dimensions['G'].width = 20
    ws_list.column_dimensions['H'].width = 22
    ws_list.column_dimensions['I'].width = 24
    ws_list.column_dimensions['J'].width = 22
    
    ws_penerimaan.column_dimensions['A'].width = 20
    ws_penerimaan.column_dimensions['B'].width = 36
    ws_penerimaan.column_dimensions['C'].width = 18
    ws_penerimaan.column_dimensions['D'].width = 16
    ws_penerimaan.column_dimensions['E'].width = 20
    ws_penerimaan.column_dimensions['F'].width = 22
    ws_penerimaan.column_dimensions['G'].width = 22
    
    ws_distribusi.column_dimensions['A'].width = 20
    ws_distribusi.column_dimensions['B'].width = 36
    ws_distribusi.column_dimensions['C'].width = 18
    ws_distribusi.column_dimensions['D'].width = 16
    ws_distribusi.column_dimensions['E'].width = 26
    ws_distribusi.column_dimensions['F'].width = 24
    ws_distribusi.column_dimensions['G'].width = 22

    ws_stok_awal.column_dimensions['A'].width = 44
    ws_stok_awal.column_dimensions['B'].width = 20
    ws_stok_awal.column_dimensions['C'].width = 36
    ws_stok_awal.column_dimensions['D'].width = 18
    ws_stok_awal.column_dimensions['E'].width = 30

    ws_stok.column_dimensions['A'].width = 20
    ws_stok.column_dimensions['B'].width = 36
    ws_stok.column_dimensions['C'].width = 18
    ws_stok.column_dimensions['D'].width = 22

    ws_kadaluwarsa.column_dimensions['A'].width = 20
    ws_kadaluwarsa.column_dimensions['B'].width = 36
    ws_kadaluwarsa.column_dimensions['C'].width = 18
    ws_kadaluwarsa.column_dimensions['D'].width = 22
    ws_kadaluwarsa.column_dimensions['E'].width = 20
    ws_kadaluwarsa.column_dimensions['F'].width = 16

    wb.save(output_path)
    print(f"File Excel Berhasil Dibuat: {output_path}")

if __name__ == "__main__":
    target = r"c:\My Data\01 AKTIF\Semester 7\KP - Capstone - PPI - TA\KP - AI Assisted Project\Dashboard E-LUNA\database_logistik_bpbd_jabar.xlsx"
    build_bpbd_excel_database(target)
